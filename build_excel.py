import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from datetime import datetime

tasks = [
    {"Task": "DB Schema 설계", "Status": "Done", "Priority": "High", "Category": "Data",
     "Due Date": "2026-09-03", "Description": "Data 영역의 High 우선순위 업무로, 2026-09-03까지 DB Schema 설계를 완료한다.",
     "Estimate": None, "Notion URL": "https://app.notion.com/p/3d536a231ca781f2ba96d59cd36533bc"},
    {"Task": "로그인 화면 개발", "Status": "Not started", "Priority": "High", "Category": "Frontend",
     "Due Date": "2026-09-04", "Description": "Frontend 영역의 High 우선순위 업무로, 2026-09-04까지 로그인 화면 개발을 완료한다.",
     "Estimate": None, "Notion URL": "https://app.notion.com/p/3d536a231ca781238d6ced018034f61b"},
    {"Task": "AI API 연동", "Status": "Not started", "Priority": "Medium", "Category": "Backend",
     "Due Date": "2026-09-09", "Description": "Backend 영역의 Medium 우선순위 업무로, 2026-09-09까지 AI API 연동을 완료한다.",
     "Estimate": None, "Notion URL": "https://app.notion.com/p/3d536a231ca7818b81cfc5b9de668359"},
    {"Task": "여행 검색 화면 개발", "Status": "Not started", "Priority": "Medium", "Category": "Frontend",
     "Due Date": "2026-09-11", "Description": "Frontend 영역의 Medium 우선순위 업무로, 2026-09-11까지 여행 검색 화면 개발을 완료한다.",
     "Estimate": None, "Notion URL": "https://app.notion.com/p/3d536a231ca7817a8b3cc8c9f6b1836a"},
    {"Task": "배포 환경 구성", "Status": "Not started", "Priority": "Low", "Category": "DevOps",
     "Due Date": "2026-09-12", "Description": "DevOps 영역의 Low 우선순위 업무로, 2026-09-12까지 배포 환경 구성을 완료한다.",
     "Estimate": None, "Notion URL": "https://app.notion.com/p/3d536a231ca781b2bd76d473dafa83e0"},
]

tasks.sort(key=lambda t: (t["Due Date"], t["Task"]))

columns = ["Task", "Status", "Priority", "Category", "Due Date", "Description", "Estimate", "Notion URL"]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Task 목록"

header_font = Font(bold=True)
header_fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")

for col_idx, col_name in enumerate(columns, start=1):
    cell = ws.cell(row=1, column=col_idx, value=col_name)
    cell.font = header_font
    cell.fill = header_fill

for row_idx, task in enumerate(tasks, start=2):
    for col_idx, col_name in enumerate(columns, start=1):
        value = task[col_name]
        if col_name == "Due Date" and value:
            value = datetime.strptime(value, "%Y-%m-%d").date()
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        if col_name == "Due Date" and value:
            cell.number_format = "yyyy-mm-dd"
        if col_name == "Notion URL" and value:
            cell.hyperlink = value
            cell.style = "Hyperlink"

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(tasks) + 1}"

for col_idx, col_name in enumerate(columns, start=1):
    max_len = max([len(col_name)] + [len(str(t[col_name])) for t in tasks])
    ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 60)

# Summary sheet
ws2 = wb.create_sheet("요약")

def count_by(field):
    counts = {}
    for t in tasks:
        counts[t[field]] = counts.get(t[field], 0) + 1
    return counts

ws2.cell(row=1, column=1, value="Status별 Task 개수").font = header_font
r = 2
for k, v in count_by("Status").items():
    ws2.cell(row=r, column=1, value=k)
    ws2.cell(row=r, column=2, value=v)
    r += 1

r += 1
ws2.cell(row=r, column=1, value="Priority별 Task 개수").font = header_font
r += 1
for k, v in count_by("Priority").items():
    ws2.cell(row=r, column=1, value=k)
    ws2.cell(row=r, column=2, value=v)
    r += 1

r += 1
ws2.cell(row=r, column=1, value="Category별 Task 개수").font = header_font
r += 1
for k, v in count_by("Category").items():
    ws2.cell(row=r, column=1, value=k)
    ws2.cell(row=r, column=2, value=v)
    r += 1

r += 1
ws2.cell(row=r, column=1, value="전체 Task 개수")
ws2.cell(row=r, column=2, value=len(tasks))
r += 1
ws2.cell(row=r, column=1, value="원본 Task DB URL")
ws2.cell(row=r, column=2, value="https://app.notion.com/p/572714600d264670aeba6091c95a9e2b")

ws2.column_dimensions["A"].width = 24
ws2.column_dimensions["B"].width = 60

out_path = "/home/af-dev/AI_SERVICE/mcp_practice/output/project_tasks.xlsx"
wb.save(out_path)
print("saved:", out_path)

# Verification
wb2 = openpyxl.load_workbook(out_path)
print("sheets:", wb2.sheetnames)
ws1 = wb2["Task 목록"]
print("data rows:", ws1.max_row - 1)
print("headers:", [c.value for c in ws1[1]])
